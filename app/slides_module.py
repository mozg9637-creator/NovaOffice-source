import os

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_FILL
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches
from PySide6.QtCore import QMarginsF, QPointF, QRectF, QSizeF, Qt, QTimer
from PySide6.QtGui import (QAction, QColor, QFont, QIcon, QImage, QKeySequence, QPageLayout,
                           QPageSize, QPainter, QPen, QPixmap)
from PySide6.QtWidgets import (QApplication, QColorDialog, QFileDialog, QHBoxLayout, QInputDialog,
                               QLabel, QListView, QListWidget, QListWidgetItem, QMessageBox,
                               QPlainTextEdit, QPushButton, QSizePolicy, QSplitter, QToolBar,
                               QVBoxLayout, QWidget)
from PySide6.QtPrintSupport import QPrintDialog, QPrinter

from common import add_recent, maybe_save

EMU_PT = 12700


# ---------------------------------------------------------------- рендер
def draw_text(p, shp, k, x, y, w, h):
    name = shp.name.lower()
    is_title = shp.is_placeholder and "title" in name and "sub" not in name
    bullets = shp.is_placeholder and "content placeholder" in name
    x += 91440 * k
    w -= 2 * 91440 * k
    yy = y + 45720 * k
    for para in shp.text_frame.paragraphs:
        runs = para.runs
        text = "".join(r.text for r in runs)
        size = next((r.font.size.pt for r in runs if r.font.size), None)
        if size is None and para.font.size:
            size = para.font.size.pt
        size = size or (40 if is_title else 20)
        px = max(6, int(size * EMU_PT * k))
        font = QFont(runs[0].font.name if runs and runs[0].font.name else "Segoe UI")
        font.setPixelSize(px)
        font.setBold(is_title or any(r.font.bold for r in runs))
        font.setItalic(any(r.font.italic for r in runs))
        font.setUnderline(any(r.font.underline for r in runs))
        color = QColor("#111111")
        for r in runs:
            try:
                if r.font.color and r.font.color.rgb is not None:
                    color = QColor("#" + str(r.font.color.rgb))
                    break
            except Exception:
                pass
        flags = Qt.TextWordWrap | Qt.AlignLeft
        if para.alignment == PP_ALIGN.CENTER:
            flags = Qt.TextWordWrap | Qt.AlignHCenter
        elif para.alignment == PP_ALIGN.RIGHT:
            flags = Qt.TextWordWrap | Qt.AlignRight
        elif is_title and "center" in name:
            flags = Qt.TextWordWrap | Qt.AlignHCenter
        if bullets and text:
            text = "•  " + text
        p.setFont(font)
        p.setPen(color)
        rect = QRectF(x, yy, max(w, 10), 100000)
        bound = p.boundingRect(rect, flags, text or " ")
        p.drawText(rect, flags, text)
        yy += bound.height() + px * 0.25


def draw_table(p, shp, k, x, y):
    t = shp.table
    ys = y
    for row in t.rows:
        rh = row.height * k
        xs = x
        for ci, cell in enumerate(row.cells):
            cw = t.columns[ci].width * k
            r = QRectF(xs, ys, cw, rh)
            p.setPen(QPen(QColor("#999999"), 1))
            p.setBrush(Qt.NoBrush)
            p.drawRect(r)
            f = QFont("Segoe UI")
            f.setPixelSize(max(6, int(14 * EMU_PT * k)))
            p.setFont(f)
            p.setPen(QColor("#111111"))
            p.drawText(r.adjusted(5, 3, -5, -3), Qt.TextWordWrap, cell.text)
            xs += cw
        ys += rh


def draw_shape(p, shp, k):
    try:
        x, y, w, h = [(getattr(shp, a) or 0) * k for a in ("left", "top", "width", "height")]
    except Exception:
        return
    try:
        st = shp.shape_type
    except Exception:
        st = None
    if st == MSO_SHAPE_TYPE.GROUP:
        for c in shp.shapes:
            draw_shape(p, c, k)
        return
    if st == MSO_SHAPE_TYPE.PICTURE:
        try:
            p.drawImage(QRectF(x, y, w, h), QImage.fromData(shp.image.blob))
        except Exception:
            pass
        return
    if getattr(shp, "has_table", False) and shp.has_table:
        draw_table(p, shp, k, x, y)
        return
    try:
        if shp.fill.type == MSO_FILL.SOLID:
            p.setPen(Qt.NoPen)
            p.setBrush(QColor("#" + str(shp.fill.fore_color.rgb)))
            try:
                oval = "OVAL" in str(shp.auto_shape_type)
            except Exception:
                oval = False
            (p.drawEllipse if oval else p.drawRect)(QRectF(x, y, w, h))
    except Exception:
        pass
    try:
        if shp.has_text_frame and shp.text_frame.text.strip():
            draw_text(p, shp, k, x, y, w, h)
    except Exception:
        pass


def render_slide(prs, idx, width=1280):
    k = width / prs.slide_width
    img = QImage(width, max(1, int(prs.slide_height * k)), QImage.Format_RGB32)
    img.fill(QColor("white"))
    slide = prs.slides[idx]
    try:
        if slide._element.cSld.bg is not None:
            bg = slide.background.fill
            if bg.type == MSO_FILL.SOLID:
                img.fill(QColor("#" + str(bg.fore_color.rgb)))
    except Exception:
        pass
    p = QPainter(img)
    p.setRenderHint(QPainter.Antialiasing)
    p.setRenderHint(QPainter.TextAntialiasing)
    p.setRenderHint(QPainter.SmoothPixmapTransform)
    for shp in slide.shapes:
        draw_shape(p, shp, k)
    p.end()
    return img


def new_prs():
    prs = Presentation()
    old = prs.slide_width
    new = Inches(13.333)
    sx = new / old
    prs.slide_width = int(new)
    for part in [prs.slide_master] + list(prs.slide_layouts):
        for sh in part.shapes:
            try:
                sh.left = int(sh.left * sx)
                sh.width = int(sh.width * sx)
            except Exception:
                pass
    s = prs.slides.add_slide(prs.slide_layouts[0])
    s.shapes.title.text = "Новая презентация"
    s.placeholders[1].text = "Подзаголовок"
    return prs


def set_text(shape, text):
    tf = shape.text_frame
    keep = None
    for para in tf.paragraphs:
        if para.runs:
            f = para.runs[0].font
            try:
                col = f.color.rgb
            except Exception:
                col = None
            keep = (f.size, f.bold, f.italic, f.name, col)
            break
    tf.text = text
    if keep:
        for para in tf.paragraphs:
            for r in para.runs:
                r.font.size, r.font.bold, r.font.italic = keep[0], keep[1], keep[2]
                if keep[3]:
                    r.font.name = keep[3]
                if keep[4] is not None:
                    r.font.color.rgb = keep[4]


# ---------------------------------------------------------------- показ / печать
class Show(QWidget):
    def __init__(self, prs, start):
        super().__init__()
        self.prs, self.i = prs, start
        self.cache = None
        self.setWindowFlag(Qt.FramelessWindowHint)
        self.setStyleSheet("background:black;")
        self.setCursor(Qt.BlankCursor)

    def go(self, n):
        n = max(0, min(len(self.prs.slides) - 1, n))
        if n != self.i or self.cache is None:
            self.i = n
            self.cache = None
        self.update()

    def paintEvent(self, e):
        if self.cache is None:
            self.cache = render_slide(self.prs, self.i, max(1280, self.width()))
        p = QPainter(self)
        p.fillRect(self.rect(), Qt.black)
        img = self.cache.scaled(self.size(), Qt.KeepAspectRatio, Qt.SmoothTransformation)
        p.drawImage((self.width() - img.width()) // 2, (self.height() - img.height()) // 2, img)
        p.end()

    def keyPressEvent(self, e):
        if e.key() in (Qt.Key_Right, Qt.Key_Down, Qt.Key_Space, Qt.Key_PageDown, Qt.Key_Return):
            self.go(self.i + 1)
        elif e.key() in (Qt.Key_Left, Qt.Key_Up, Qt.Key_PageUp, Qt.Key_Backspace):
            self.go(self.i - 1)
        elif e.key() == Qt.Key_Escape:
            self.close()

    def mousePressEvent(self, e):
        self.go(self.i + (1 if e.button() == Qt.LeftButton else -1))


# ---------------------------------------------------------------- модуль
class SlidesModule(QWidget):
    def __init__(self):
        super().__init__()
        self.prs = new_prs()
        self.path = None
        self.idx = 0
        self.dirty = False
        self.cur_img = None
        self.shapes = []
        self.loading = False
        self.show_win = None

        root = QVBoxLayout(self)
        root.setContentsMargins(0, 0, 0, 0)
        root.setSpacing(0)
        bar = QToolBar()
        bar.setMovable(False)
        root.addWidget(bar)
        self.act(bar, "🆕 Новая", "Новая презентация", self.new)
        self.act(bar, "📂 Открыть", "Открыть .pptx", self.open_dialog, "Ctrl+O")
        self.act(bar, "💾 Сохранить", "Сохранить", self.save, "Ctrl+S")
        self.act(bar, "💾 Как…", "Сохранить как…", self.save_as, "Ctrl+Shift+S")
        bar.addSeparator()
        self.act(bar, "➕ Слайд", "Добавить слайд", self.add_slide)
        self.act(bar, "🗑 Слайд", "Удалить слайд", self.del_slide)
        self.act(bar, "⬆", "Переместить вверх", lambda: self.move_slide(-1))
        self.act(bar, "⬇", "Переместить вниз", lambda: self.move_slide(1))
        bar.addSeparator()
        self.act(bar, "🔤 Текст", "Добавить текстовый блок", self.add_textbox)
        self.act(bar, "🖼 Картинка", "Вставить картинку", self.add_picture)
        self.act(bar, "🎨 Фон", "Цвет фона слайда", self.set_bg)
        bar.addSeparator()
        self.act(bar, "▶ Показ", "Показ с текущего слайда", self.present, "F5")
        self.act(bar, "📄 В PDF", "Экспорт в PDF", self.export_pdf)
        self.act(bar, "🖨 Печать", "Печать", self.print_slides, "Ctrl+P")

        self.list = QListWidget()
        self.list.setViewMode(QListView.ListMode)
        self.list.setIconSize(QListWidget().iconSize().expandedTo(self.list.iconSize()))
        from PySide6.QtCore import QSize
        self.list.setIconSize(QSize(180, 101))
        self.list.setFixedWidth(215)
        self.list.setSpacing(5)
        self.list.currentRowChanged.connect(self.select_slide)

        self.canvas = QLabel()
        self.canvas.setAlignment(Qt.AlignCenter)
        self.canvas.setSizePolicy(QSizePolicy.Ignored, QSizePolicy.Ignored)
        self.canvas.setMinimumSize(300, 200)
        self.canvas.setStyleSheet("background:#0f1117;")

        side = QWidget()
        side.setFixedWidth(290)
        sl = QVBoxLayout(side)
        sl.setContentsMargins(12, 12, 12, 12)
        sl.addWidget(QLabel("Элементы слайда"))
        self.shape_list = QListWidget()
        self.shape_list.setMaximumHeight(180)
        self.shape_list.currentRowChanged.connect(self.select_shape)
        sl.addWidget(self.shape_list)
        self.editor = QPlainTextEdit()
        self.editor.setPlaceholderText("Выберите элемент и редактируйте текст — слайд обновится сразу")
        self.editor.textChanged.connect(lambda: self.timer.start(350))
        sl.addWidget(self.editor, 1)
        b = QPushButton("🗑 Удалить элемент")
        b.clicked.connect(self.del_shape)
        sl.addWidget(b)
        self.timer = QTimer(self)
        self.timer.setSingleShot(True)
        self.timer.timeout.connect(self.apply_text)

        split = QSplitter()
        split.addWidget(self.list)
        split.addWidget(self.canvas)
        split.addWidget(side)
        split.setStretchFactor(1, 1)
        root.addWidget(split, 1)
        self.status = QLabel("")
        self.status.setObjectName("status")
        self.status.setContentsMargins(12, 4, 12, 4)
        root.addWidget(self.status)
        self.reload_all()

    def act(self, bar, text, tip, slot, key=None):
        a = QAction(text, self)
        a.setToolTip(tip + (f" ({key})" if key else ""))
        if key:
            a.setShortcut(QKeySequence(key))
        a.triggered.connect(lambda *_: slot())
        bar.addAction(a)

    # ---- файл
    def can_close(self):
        return maybe_save(self, self.dirty, os.path.basename(self.path or "презентация.pptx"), self.save)

    def new(self):
        if not self.can_close():
            return
        self.prs, self.path, self.dirty, self.idx = new_prs(), None, False, 0
        self.reload_all()

    def open_dialog(self):
        p, _ = QFileDialog.getOpenFileName(self, "Открыть презентацию", "", "PowerPoint (*.pptx)")
        if p:
            self.open_file(p)

    def open_file(self, path):
        if not self.can_close():
            return
        try:
            prs = Presentation(path)
            if len(prs.slides) == 0:
                prs.slides.add_slide(prs.slide_layouts[0])
        except Exception as e:
            QMessageBox.warning(self, "Презентация", f"Не удалось открыть файл:\n{e}")
            return
        self.prs, self.path, self.dirty, self.idx = prs, path, False, 0
        add_recent(path)
        self.reload_all()

    def save(self):
        return self.save_as() if not self.path else self._write(self.path)

    def save_as(self):
        p, _ = QFileDialog.getSaveFileName(self, "Сохранить презентацию", self.path or "",
                                           "PowerPoint (*.pptx)")
        if not p:
            return False
        if not p.lower().endswith(".pptx"):
            p += ".pptx"
        return self._write(p)

    def _write(self, path):
        try:
            self.prs.save(path)
        except Exception as e:
            QMessageBox.critical(self, "Презентация", f"Ошибка сохранения:\n{e}")
            return False
        self.path, self.dirty = path, False
        add_recent(path)
        self.status.setText("Сохранено: " + path)
        return True

    # ---- отображение
    def slide(self):
        return self.prs.slides[self.idx]

    def reload_all(self):
        QApplication.setOverrideCursor(Qt.WaitCursor)
        try:
            self.list.blockSignals(True)
            self.list.clear()
            for i in range(len(self.prs.slides)):
                it = QListWidgetItem(self.thumb(i), str(i + 1))
                self.list.addItem(it)
            self.idx = max(0, min(self.idx, len(self.prs.slides) - 1))
            self.list.setCurrentRow(self.idx)
            self.list.blockSignals(False)
            self.refresh()
        finally:
            QApplication.restoreOverrideCursor()

    def thumb(self, i):
        return QIcon(QPixmap.fromImage(render_slide(self.prs, i, 360)))

    def refresh(self, thumb=False):
        self.cur_img = render_slide(self.prs, self.idx, 1600)
        self.fit()
        if thumb:
            it = self.list.item(self.idx)
            if it:
                it.setIcon(self.thumb(self.idx))
        self.fill_shapes()
        self.status.setText(f"Слайд {self.idx + 1} из {len(self.prs.slides)}")

    def fit(self):
        if self.cur_img is None:
            return
        pm = QPixmap.fromImage(self.cur_img).scaled(self.canvas.size() - self.canvas.size() / 20,
                                                    Qt.KeepAspectRatio, Qt.SmoothTransformation)
        self.canvas.setPixmap(pm)

    def resizeEvent(self, e):
        super().resizeEvent(e)
        self.fit()

    def showEvent(self, e):
        super().showEvent(e)
        QTimer.singleShot(0, self.fit)

    def select_slide(self, n):
        if n >= 0:
            self.idx = n
            self.refresh()

    # ---- элементы
    def fill_shapes(self, keep=-1):
        self.loading = True
        self.shapes = [s for s in self.slide().shapes if getattr(s, "has_text_frame", False)]
        self.shape_list.blockSignals(True)
        self.shape_list.clear()
        for s in self.shapes:
            t = s.text_frame.text.strip().replace("\n", " ")
            self.shape_list.addItem(f"{s.name}: {t[:28]}" if t else f"{s.name} (пусто)")
        self.shape_list.blockSignals(False)
        self.editor.setPlainText("")
        self.loading = False
        if self.shapes:
            self.shape_list.setCurrentRow(keep if 0 <= keep < len(self.shapes) else 0)

    def select_shape(self, n):
        self.loading = True
        self.editor.setPlainText(self.shapes[n].text_frame.text if 0 <= n < len(self.shapes) else "")
        self.loading = False

    def apply_text(self):
        n = self.shape_list.currentRow()
        if self.loading or not (0 <= n < len(self.shapes)):
            return
        set_text(self.shapes[n], self.editor.toPlainText())
        self.dirty = True
        self.cur_img = render_slide(self.prs, self.idx, 1600)
        self.fit()
        self.list.item(self.idx).setIcon(self.thumb(self.idx))

    def del_shape(self):
        n = self.shape_list.currentRow()
        if 0 <= n < len(self.shapes):
            el = self.shapes[n]._element
            el.getparent().remove(el)
            self.dirty = True
            self.refresh(thumb=True)

    # ---- операции со слайдами
    def add_slide(self):
        names = [l.name for l in self.prs.slide_layouts]
        name, ok = QInputDialog.getItem(self, "Макет", "Выберите макет слайда:", names,
                                        min(1, len(names) - 1), False)
        if not ok:
            return
        self.prs.slides.add_slide(self.prs.slide_layouts[names.index(name)])
        self.dirty = True
        self.idx = len(self.prs.slides) - 1
        self.reload_all()

    def del_slide(self):
        if len(self.prs.slides) < 2:
            return
        lst = self.prs.slides._sldIdLst
        sld = list(lst)[self.idx]
        self.prs.part.drop_rel(sld.rId)
        lst.remove(sld)
        self.dirty = True
        self.idx = max(0, self.idx - 1)
        self.reload_all()

    def move_slide(self, d):
        n = self.idx + d
        if not (0 <= n < len(self.prs.slides)):
            return
        lst = self.prs.slides._sldIdLst
        items = list(lst)
        el = items[self.idx]
        lst.remove(el)
        lst.insert(n, el)
        self.idx = n
        self.dirty = True
        self.reload_all()

    def add_textbox(self):
        sw, sh = self.prs.slide_width, self.prs.slide_height
        tb = self.slide().shapes.add_textbox(int(sw * 0.1), int(sh * 0.4), int(sw * 0.8), int(sh * 0.2))
        tb.text_frame.text = "Новый текст"
        tb.text_frame.paragraphs[0].runs[0].font.size = Inches(0.33)
        self.dirty = True
        self.refresh(thumb=True)
        self.shape_list.setCurrentRow(len(self.shapes) - 1)

    def add_picture(self):
        p, _ = QFileDialog.getOpenFileName(self, "Картинка", "", "Изображения (*.png *.jpg *.jpeg *.bmp *.gif)")
        if not p:
            return
        sw, sh = self.prs.slide_width, self.prs.slide_height
        try:
            pic = self.slide().shapes.add_picture(p, int(sw * 0.25), int(sh * 0.25), width=int(sw * 0.5))
            if pic.height > sh * 0.9:
                ratio = sh * 0.9 / pic.height
                pic.height, pic.width = int(pic.height * ratio), int(pic.width * ratio)
        except Exception as e:
            QMessageBox.warning(self, "Картинка", str(e))
            return
        self.dirty = True
        self.refresh(thumb=True)

    def set_bg(self):
        c = QColorDialog.getColor(QColor("white"), self, "Фон слайда")
        if c.isValid():
            f = self.slide().background.fill
            f.solid()
            f.fore_color.rgb = RGBColor(c.red(), c.green(), c.blue())
            self.dirty = True
            self.refresh(thumb=True)

    # ---- показ / PDF / печать
    def present(self):
        self.show_win = Show(self.prs, self.idx)
        self.show_win.showFullScreen()
        self.show_win.setFocus()

    def paint_slides(self, pr, a, b):
        painter = QPainter(pr)
        r = pr.pageRect(QPrinter.DevicePixel)
        for i in range(a, b + 1):
            if i > a:
                pr.newPage()
            img = render_slide(self.prs, i, 2000).scaled(int(r.width()), int(r.height()),
                                                         Qt.KeepAspectRatio, Qt.SmoothTransformation)
            painter.drawImage(QPointF((r.width() - img.width()) / 2, (r.height() - img.height()) / 2), img)
        painter.end()

    def export_pdf(self):
        p, _ = QFileDialog.getSaveFileName(self, "Экспорт в PDF", "", "PDF (*.pdf)")
        if not p:
            return
        pr = QPrinter(QPrinter.HighResolution)
        pr.setOutputFormat(QPrinter.PdfFormat)
        pr.setOutputFileName(p)
        w, h = self.prs.slide_width / 36000, self.prs.slide_height / 36000
        pr.setPageSize(QPageSize(QSizeF(w, h), QPageSize.Millimeter))
        pr.setPageMargins(QMarginsF(0, 0, 0, 0), QPageLayout.Millimeter)
        self.paint_slides(pr, 0, len(self.prs.slides) - 1)
        self.status.setText("PDF сохранён: " + p)

    def print_slides(self):
        pr = QPrinter(QPrinter.HighResolution)
        pr.setPageOrientation(QPageLayout.Landscape)
        dlg = QPrintDialog(pr, self)
        dlg.setMinMax(1, len(self.prs.slides))
        dlg.setOption(QPrintDialog.PrintPageRange, True)
        if dlg.exec() != QPrintDialog.Accepted:
            return
        n = len(self.prs.slides)
        if pr.printRange() == QPrinter.PageRange:
            a, b = max(0, pr.fromPage() - 1), min(n - 1, pr.toPage() - 1)
        else:
            a, b = 0, n - 1
        self.paint_slides(pr, a, b)
