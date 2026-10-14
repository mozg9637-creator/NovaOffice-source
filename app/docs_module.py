import io
import os
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from docx.table import Table
from docx.text.paragraph import Paragraph
from PySide6.QtCore import QByteArray, QBuffer, QIODevice, QMarginsF, QSizeF, Qt, QUrl
from PySide6.QtGui import (QAction, QColor, QFont, QImage, QKeySequence, QPageLayout, QPageSize,
                           QTextBlockFormat, QTextCharFormat, QTextCursor, QTextDocument,
                           QTextDocumentWriter, QTextListFormat, QTextTableFormat)
from PySide6.QtWidgets import (QColorDialog, QComboBox, QFileDialog, QFontComboBox, QHBoxLayout,
                               QInputDialog, QLabel, QMessageBox, QTextEdit, QToolBar, QVBoxLayout,
                               QWidget)
from PySide6.QtPrintSupport import QPageSetupDialog, QPrintDialog, QPrintPreviewDialog, QPrinter

from common import add_recent, maybe_save

W_NS = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"
HEAD_PT = {1: 22, 2: 17, 3: 14}
CTRL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def family_of(cf):
    try:
        f = cf.fontFamilies()
        if isinstance(f, (list, tuple)) and f:
            return str(f[0])
    except Exception:
        pass
    try:
        return cf.fontFamily() or None
    except Exception:
        return None


# ------------------------------------------------------------------ DOCX чтение
def load_docx(edit, path):
    d = Document(path)
    qdoc = edit.document()
    edit.clear()
    cur = edit.textCursor()
    first = True
    n_img = 0
    for child in d.element.body.iterchildren():
        tag = child.tag.split("}")[-1]
        if tag == "p":
            para = Paragraph(child, d)
            bf = QTextBlockFormat()
            try:
                v = int(para.alignment) if para.alignment is not None else 0
            except Exception:
                v = 0
            bf.setAlignment({1: Qt.AlignHCenter, 2: Qt.AlignRight, 3: Qt.AlignJustify}.get(v, Qt.AlignLeft))
            sname = para.style.name if para.style is not None else ""
            last = sname.split()[-1] if sname else ""
            level = int(last) if sname.startswith("Heading") and last.isdigit() else 0
            if level:
                bf.setHeadingLevel(min(level, 6))
            if first:
                cur.setBlockFormat(bf)
                first = False
            else:
                cur.insertBlock(bf, QTextCharFormat())
            if sname.startswith("List"):
                cur.insertText("•  ")
            for run in para.runs:
                cf = QTextCharFormat()
                f = run.font
                if run.bold or level:
                    cf.setFontWeight(QFont.Bold)
                if run.italic:
                    cf.setFontItalic(True)
                if run.underline:
                    cf.setFontUnderline(True)
                if f.size:
                    cf.setFontPointSize(f.size.pt)
                elif level:
                    cf.setFontPointSize(HEAD_PT.get(level, 13))
                if f.name:
                    cf.setFontFamilies([f.name])
                try:
                    if f.color is not None and f.color.rgb is not None:
                        cf.setForeground(QColor("#" + str(f.color.rgb)))
                except Exception:
                    pass
                try:
                    for blip in run._element.xpath(".//a:blip"):
                        blob = d.part.related_parts[blip.get(W_NS)].blob
                        img = QImage.fromData(blob)
                        if not img.isNull():
                            n_img += 1
                            name = f"docx_img_{n_img}"
                            qdoc.addResource(QTextDocument.ImageResource, QUrl(name), img)
                            from PySide6.QtGui import QTextImageFormat
                            imf = QTextImageFormat()
                            imf.setName(name)
                            imf.setWidth(min(img.width(), 680))
                            cur.insertImage(imf)
                except Exception:
                    pass
                if run.text:
                    cur.insertText(run.text, cf)
        elif tag == "tbl":
            t = Table(child, d)
            rows, cols = len(t.rows), len(t.columns)
            if rows and cols:
                fmt = QTextTableFormat()
                fmt.setBorder(1)
                fmt.setCellPadding(5)
                fmt.setBorderCollapse(True)
                table = cur.insertTable(rows, cols, fmt)
                for r in range(rows):
                    for c in range(cols):
                        try:
                            table.cellAt(r, c).firstCursorPosition().insertText(t.cell(r, c).text)
                        except Exception:
                            pass
                cur.movePosition(QTextCursor.End)
                first = True
    edit.document().setModified(False)


# ------------------------------------------------------------------ DOCX запись
def runs_of(qdoc, block):
    """Разбивает блок на (текст, формат) без QTextBlock.iterator."""
    start, n = block.position(), block.length() - 1
    cur = QTextCursor(qdoc)
    out, text, cf0 = [], "", None
    for i in range(n):
        cur.setPosition(start + i + 1)
        cf = cur.charFormat()
        ch = qdoc.characterAt(start + i)
        if cf0 is not None and cf == cf0 and not cf.isImageFormat():
            text += ch
        else:
            if cf0 is not None:
                out.append((text, cf0))
            cf0, text = cf, ch
    if cf0 is not None:
        out.append((text, cf0))
    return out


def export_docx(qdoc, path):
    d = Document()
    done_tables = set()
    block = qdoc.begin()
    while block.isValid():
        t = QTextCursor(block).currentTable()
        if t is not None:
            key = t.firstPosition()
            if key not in done_tables:
                done_tables.add(key)
                table = d.add_table(rows=t.rows(), cols=t.columns())
                table.style = "Table Grid"
                for r in range(t.rows()):
                    for c in range(t.columns()):
                        cell = t.cellAt(r, c)
                        cc = cell.firstCursorPosition()
                        cc.setPosition(cell.lastCursorPosition().position(), QTextCursor.KeepAnchor)
                        table.cell(r, c).text = CTRL.sub("", cc.selectedText().replace("\u2029", "\n"))
            block = block.next()
            continue
        bf = block.blockFormat()
        style = None
        lst = block.textList()
        if lst is not None:
            style = "List Number" if lst.format().style() == QTextListFormat.ListDecimal else "List Bullet"
        elif bf.headingLevel() > 0:
            style = f"Heading {min(bf.headingLevel(), 9)}"
        try:
            p = d.add_paragraph(style=style)
        except KeyError:
            p = d.add_paragraph()
        al = bf.alignment()
        if al & Qt.AlignHCenter:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif al & Qt.AlignRight:
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        elif al & Qt.AlignJustify:
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        for text, cf in runs_of(qdoc, block):
            if cf.isImageFormat():
                try:
                    imf = cf.toImageFormat()
                    img = qdoc.resource(QTextDocument.ImageResource, QUrl(imf.name()))
                    if isinstance(img, QImage) and not img.isNull():
                        ba = QByteArray()
                        buf = QBuffer(ba)
                        buf.open(QIODevice.WriteOnly)
                        img.save(buf, "PNG")
                        p.add_run().add_picture(io.BytesIO(bytes(ba)), width=Inches(min(6.0, img.width() / 96)))
                except Exception:
                    pass
                continue
            r = p.add_run(CTRL.sub("", text.replace("\u2028", "\n")))
            r.bold = (cf.fontWeight() >= 600) or None
            r.italic = cf.fontItalic() or None
            r.underline = cf.fontUnderline() or None
            if cf.fontPointSize() > 0:
                r.font.size = Pt(cf.fontPointSize())
            fam = family_of(cf)
            if fam:
                r.font.name = fam
            br = cf.foreground()
            if br.style() != Qt.NoBrush:
                c = br.color()
                r.font.color.rgb = RGBColor(c.red(), c.green(), c.blue())
        block = block.next()
    d.save(path)


class DocsModule(QWidget):
    def __init__(self):
        super().__init__()
        self.path = None
        self.printer = QPrinter(QPrinter.HighResolution)
        self.printer.setPageSize(QPageSize(QPageSize.A4))
        self.printer.setPageMargins(QMarginsF(20, 20, 20, 20), QPageLayout.Millimeter)

        root = QVBoxLayout(self)
        root.setContentsMargins(0, 0, 0, 0)
        root.setSpacing(0)

        bar = QToolBar()
        bar.setMovable(False)
        root.addWidget(bar)
        self.act(bar, "Новый", "Новый документ", self.new, "Ctrl+N")
        self.act(bar, "Открыть", "Открыть .docx / .md / .txt / .html", self.open_dialog, "Ctrl+O")
        self.act(bar, "Сохранить", "Сохранить", self.save, "Ctrl+S")
        self.act(bar, "Сохранить как…", "Сохранить как…", self.save_as, "Ctrl+Shift+S")
        bar.addSeparator()
        self.act(bar, "В PDF", "Экспорт в PDF", self.export_pdf)
        self.act(bar, "Просмотр", "Предварительный просмотр печати", self.preview)
        self.act(bar, "Печать", "Печать", self.print_doc, "Ctrl+P")
        self.act(bar, "Страница", "Размер бумаги, поля, ориентация", self.page_setup)
        bar.addSeparator()
        self.act(bar, "Найти", "Найти в тексте", self.find, "Ctrl+F")
        self.act(bar, "Заменить", "Найти и заменить всё", self.replace_all, "Ctrl+H")
        self.act(bar, "Дата", "Вставить сегодняшнюю дату", self.insert_date)

        fb = QToolBar()
        fb.setMovable(False)
        root.addWidget(fb)
        self.style_box = QComboBox()
        self.style_box.addItems(["Обычный", "Заголовок 1", "Заголовок 2", "Заголовок 3"])
        self.style_box.activated.connect(self.set_style)
        fb.addWidget(self.style_box)
        self.font_box = QFontComboBox()
        self.font_box.setMaximumWidth(190)
        self.font_box.activated.connect(self.set_font)
        fb.addWidget(self.font_box)
        self.size_box = QComboBox()
        self.size_box.setEditable(True)
        self.size_box.addItems([str(s) for s in (8, 9, 10, 11, 12, 14, 16, 18, 20, 24, 28, 36, 48, 72)])
        self.size_box.setCurrentText("11")
        self.size_box.setMaximumWidth(70)
        self.size_box.activated.connect(self.set_size)
        self.size_box.lineEdit().editingFinished.connect(self.set_size)
        fb.addWidget(self.size_box)
        fb.addSeparator()
        self.a_bold = self.act(fb, "𝐁", "Жирный", self.fmt_bold, "Ctrl+B", True)
        self.a_ital = self.act(fb, "𝑰", "Курсив", self.fmt_ital, "Ctrl+I", True)
        self.a_under = self.act(fb, "U̲", "Подчёркнутый", self.fmt_under, "Ctrl+U", True)
        self.act(fb, "Цвет", "Цвет текста", self.fmt_color)
        fb.addSeparator()
        self.act(fb, "Слева", "По левому краю", lambda: self.edit.setAlignment(Qt.AlignLeft))
        self.act(fb, "Центр", "По центру", lambda: self.edit.setAlignment(Qt.AlignHCenter))
        self.act(fb, "Справа", "По правому краю", lambda: self.edit.setAlignment(Qt.AlignRight))
        self.act(fb, "По ширине", "По ширине", lambda: self.edit.setAlignment(Qt.AlignJustify))
        fb.addSeparator()
        self.act(fb, "• Список", "Маркированный список", lambda: self.toggle_list(QTextListFormat.ListDisc))
        self.act(fb, "1. Список", "Нумерованный список", lambda: self.toggle_list(QTextListFormat.ListDecimal))
        self.act(fb, "Картинка", "Вставить картинку", self.insert_image)
        self.act(fb, "Таблица", "Вставить таблицу", self.insert_table)
        fb.addSeparator()
        self.act(fb, "Отменить", "Отменить", lambda: self.edit.undo())
        self.act(fb, "Повторить", "Повторить", lambda: self.edit.redo())

        self.edit = QTextEdit()
        self.edit.setAcceptRichText(True)
        self.edit.setFrameShape(QTextEdit.NoFrame)
        self.edit.document().setDefaultFont(QFont("Segoe UI", 11))
        self.edit.document().setDocumentMargin(56)
        self.edit.setStyleSheet("QTextEdit{background:#ffffff;color:#16181d;"
                                "selection-background-color:#a9bbff;selection-color:#000;}")
        self.edit.setMaximumWidth(900)
        self.edit.cursorPositionChanged.connect(self.sync_toolbar)
        self.edit.textChanged.connect(self.update_status)

        wrap = QWidget()
        wrap.setObjectName("wrap")
        wrap.setAttribute(Qt.WA_StyledBackground, True)
        hl = QHBoxLayout(wrap)
        hl.setContentsMargins(0, 14, 0, 14)
        hl.addStretch(1)
        hl.addWidget(self.edit, 100)
        hl.addStretch(1)
        root.addWidget(wrap, 1)

        self.status = QLabel("")
        self.status.setObjectName("status")
        self.status.setContentsMargins(12, 4, 12, 4)
        root.addWidget(self.status)
        self.update_status()

    def act(self, bar, text, tip, slot, key=None, checkable=False):
        a = QAction(text, self)
        a.setToolTip(tip + (f" ({key})" if key else ""))
        if key:
            a.setShortcut(QKeySequence(key))
        a.setCheckable(checkable)
        a.triggered.connect(lambda *_: slot())
        bar.addAction(a)
        return a

    # ---- форматирование
    def merge(self, fmt):
        self.edit.mergeCurrentCharFormat(fmt)
        self.edit.setFocus()

    def fmt_bold(self):
        f = QTextCharFormat()
        f.setFontWeight(QFont.Bold if self.a_bold.isChecked() else QFont.Normal)
        self.merge(f)

    def fmt_ital(self):
        f = QTextCharFormat()
        f.setFontItalic(self.a_ital.isChecked())
        self.merge(f)

    def fmt_under(self):
        f = QTextCharFormat()
        f.setFontUnderline(self.a_under.isChecked())
        self.merge(f)

    def fmt_color(self):
        c = QColorDialog.getColor(QColor("#000000"), self, "Цвет текста")
        if c.isValid():
            f = QTextCharFormat()
            f.setForeground(c)
            self.merge(f)

    def set_font(self, *_):
        f = QTextCharFormat()
        f.setFontFamilies([self.font_box.currentFont().family()])
        self.merge(f)

    def set_size(self, *_):
        try:
            s = float(self.size_box.currentText().replace(",", "."))
        except ValueError:
            return
        if 1 <= s <= 400:
            f = QTextCharFormat()
            f.setFontPointSize(s)
            self.merge(f)

    def set_style(self, i):
        c = self.edit.textCursor()
        c.beginEditBlock()
        bf = QTextBlockFormat()
        bf.setHeadingLevel(i)
        cf = QTextCharFormat()
        cf.setFontPointSize(HEAD_PT.get(i, 11))
        cf.setFontWeight(QFont.Bold if i else QFont.Normal)
        c2 = QTextCursor(c)
        c2.movePosition(QTextCursor.StartOfBlock)
        c2.movePosition(QTextCursor.EndOfBlock, QTextCursor.KeepAnchor)
        c2.mergeCharFormat(cf)
        c.mergeBlockFormat(bf)
        c.endEditBlock()
        self.edit.setFocus()

    def toggle_list(self, style):
        c = self.edit.textCursor()
        lst = c.currentList()
        if lst is not None and lst.format().style() == style:
            c.beginEditBlock()
            lst.remove(c.block())
            bf = c.blockFormat()
            bf.setIndent(0)
            c.setBlockFormat(bf)
            c.endEditBlock()
        else:
            c.createList(style)
        self.edit.setFocus()

    def insert_image(self):
        p, _ = QFileDialog.getOpenFileName(self, "Картинка", "", "Изображения (*.png *.jpg *.jpeg *.bmp *.gif)")
        if not p:
            return
        img = QImage(p)
        if img.isNull():
            return
        from PySide6.QtGui import QTextImageFormat
        name = f"img_{id(img)}"
        self.edit.document().addResource(QTextDocument.ImageResource, QUrl(name), img)
        f = QTextImageFormat()
        f.setName(name)
        f.setWidth(min(img.width(), 680))
        self.edit.textCursor().insertImage(f)

    def insert_table(self):
        r, ok = QInputDialog.getInt(self, "Таблица", "Строк:", 3, 1, 100)
        if not ok:
            return
        c, ok = QInputDialog.getInt(self, "Таблица", "Столбцов:", 3, 1, 20)
        if not ok:
            return
        f = QTextTableFormat()
        f.setBorder(1)
        f.setCellPadding(5)
        f.setBorderCollapse(True)
        self.edit.textCursor().insertTable(r, c, f)

    def find(self):
        t, ok = QInputDialog.getText(self, "Найти", "Текст:")
        if ok and t:
            if not self.edit.find(t):
                self.edit.moveCursor(QTextCursor.Start)
                if not self.edit.find(t):
                    QMessageBox.information(self, "Поиск", "Ничего не найдено.")

    def replace_all(self):
        old, ok = QInputDialog.getText(self, "Заменить", "Найти:")
        if not ok or not old:
            return
        new, ok = QInputDialog.getText(self, "Заменить", f"Заменить «{old}» на:")
        if not ok:
            return
        doc = self.edit.document()
        cur = doc.find(old)
        n = 0
        self.edit.textCursor().beginEditBlock()
        while not cur.isNull():
            cur.insertText(new)
            n += 1
            cur = doc.find(old, cur)
        self.edit.textCursor().endEditBlock()
        self.status.setText(f"Заменено: {n}")

    def insert_date(self):
        from datetime import date
        self.edit.textCursor().insertText(date.today().strftime("%d.%m.%Y"))

    def sync_toolbar(self):
        cf = self.edit.currentCharFormat()
        self.a_bold.setChecked(cf.fontWeight() >= 600)
        self.a_ital.setChecked(cf.fontItalic())
        self.a_under.setChecked(cf.fontUnderline())
        pt = cf.fontPointSize()
        self.size_box.setCurrentText(str(int(pt)) if pt > 0 else "11")
        self.style_box.setCurrentIndex(min(3, max(0, self.edit.textCursor().blockFormat().headingLevel())))

    def update_status(self):
        t = self.edit.toPlainText()
        self.status.setText(f"Слов: {len(t.split())}   Символов: {len(t)}"
                            + (f"   ·   {self.path}" if self.path else ""))

    # ---- файлы
    def can_close(self):
        return maybe_save(self, self.edit.document().isModified(),
                          os.path.basename(self.path or "документ"), self.save)

    def new(self):
        if not self.can_close():
            return
        self.edit.clear()
        self.edit.document().setModified(False)
        self.path = None
        self.update_status()

    def open_dialog(self):
        p, _ = QFileDialog.getOpenFileName(self, "Открыть документ", "",
                                           "Документы (*.docx *.md *.txt *.html *.htm)")
        if p:
            self.open_file(p)

    def open_file(self, path):
        if not self.can_close():
            return
        ext = os.path.splitext(path)[1].lower()
        try:
            if ext == ".docx":
                load_docx(self.edit, path)
            else:
                with open(path, "rb") as f:
                    raw = f.read()
                try:
                    s = raw.decode("utf-8-sig")
                except UnicodeDecodeError:
                    s = raw.decode("cp1251", errors="replace")
                if ext == ".md":
                    self.edit.setMarkdown(s)
                elif ext in (".html", ".htm"):
                    self.edit.setHtml(s)
                else:
                    self.edit.setPlainText(s)
        except Exception as e:
            QMessageBox.warning(self, "Документ", f"Не удалось открыть файл:\n{e}")
            return
        self.edit.document().setModified(False)
        self.path = path
        add_recent(path)
        self.update_status()

    def save(self):
        if not self.path or os.path.splitext(self.path)[1].lower() not in (".docx", ".odt", ".md", ".txt", ".html", ".htm"):
            return self.save_as()
        return self._write(self.path)

    def save_as(self):
        filters = "Word (*.docx);;OpenDocument (*.odt);;Markdown (*.md);;HTML (*.html);;Текст (*.txt)"
        p, flt = QFileDialog.getSaveFileName(self, "Сохранить документ", self.path or "", filters)
        if not p:
            return False
        if not os.path.splitext(p)[1]:
            p += {"Word": ".docx", "OpenDocument": ".odt", "Markdown": ".md", "HTML": ".html"}.get(
                flt.split(" ")[0], ".txt")
        return self._write(p)

    def _write(self, path):
        qdoc = self.edit.document()
        ext = os.path.splitext(path)[1].lower()
        try:
            if ext == ".docx":
                export_docx(qdoc, path)
            elif ext == ".odt":
                w = QTextDocumentWriter(path, QByteArray(b"odf"))
                if not w.write(qdoc):
                    raise OSError("не удалось записать ODT")
            elif ext == ".md":
                with open(path, "w", encoding="utf-8") as f:
                    f.write(qdoc.toMarkdown())
            elif ext in (".html", ".htm"):
                with open(path, "w", encoding="utf-8") as f:
                    f.write(qdoc.toHtml())
            else:
                with open(path, "w", encoding="utf-8") as f:
                    f.write(qdoc.toPlainText())
        except Exception as e:
            QMessageBox.critical(self, "Документ", f"Ошибка сохранения:\n{e}")
            return False
        self.path = path
        qdoc.setModified(False)
        add_recent(path)
        self.update_status()
        return True

    # ---- печать
    def export_pdf(self):
        p, _ = QFileDialog.getSaveFileName(self, "Экспорт в PDF", "", "PDF (*.pdf)")
        if not p:
            return
        pr = QPrinter(QPrinter.HighResolution)
        pr.setOutputFormat(QPrinter.PdfFormat)
        pr.setOutputFileName(p)
        pr.setPageLayout(self.printer.pageLayout())
        self.edit.document().print_(pr)
        self.status.setText("PDF сохранён: " + p)

    def print_doc(self):
        dlg = QPrintDialog(self.printer, self)
        if dlg.exec() == QPrintDialog.Accepted:
            self.edit.document().print_(self.printer)

    def preview(self):
        dlg = QPrintPreviewDialog(self.printer, self)
        dlg.resize(1000, 800)
        dlg.paintRequested.connect(lambda pr: self.edit.document().print_(pr))
        dlg.exec()

    def page_setup(self):
        QPageSetupDialog(self.printer, self).exec()
